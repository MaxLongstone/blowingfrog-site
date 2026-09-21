#!/usr/bin/env python3
"""Turn one of the project's markdown docs into a Word document.

    python3 tools/md_to_docx.py docs/asset-build-sheet.md "/path/out.docx"

Handles headings, paragraphs, bold and inline code, lists, tables, and fenced code blocks
(shown in a shaded box so a prompt can be copied whole).
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def add_inline(par, text, size=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x8B, 0x2E, 0x1E)
        else:
            r = par.add_run(part)
        if size and not (part.startswith("`")):
            r.font.size = Pt(size)


def code_box(doc, lines):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    shade(c, "F1F1EC")
    c.paragraphs[0].text = ""
    first = True
    for ln in lines:
        p = c.paragraphs[0] if first else c.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(ln); r.font.name = "Consolas"; r.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build(md, out):
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.85)
        s.top_margin = s.bottom_margin = Inches(0.8)
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("```"):
            j = i + 1; body = []
            while j < len(lines) and not lines[j].startswith("```"):
                body.append(lines[j]); j += 1
            code_box(doc, body); i = j + 1; continue
        m = re.match(r"^(#{1,4})\s+(.*)$", ln)
        if m:
            doc.add_heading(m.group(2).strip(), level=min(len(m.group(1)), 3)); i += 1; continue
        if ln.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s\-:|]+\|$", lines[i + 1]):
            head = [c.strip() for c in ln.strip("|").split("|")]
            rows = []; j = i + 2
            while j < len(lines) and lines[j].startswith("|"):
                rows.append([c.strip() for c in lines[j].strip().strip("|").split("|")]); j += 1
            t = doc.add_table(rows=1, cols=len(head)); t.style = "Table Grid"
            for k, h in enumerate(head):
                cell = t.rows[0].cells[k]; shade(cell, "E4E7D0"); cell.paragraphs[0].text = ""
                add_inline(cell.paragraphs[0], f"**{h}**", 10)
            for r in rows:
                cells = t.add_row().cells
                for k in range(len(head)):
                    cells[k].paragraphs[0].text = ""
                    add_inline(cells[k].paragraphs[0], r[k] if k < len(r) else "", 10)
            doc.add_paragraph(); i = j; continue
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", ln)
        if m:
            style = "List Number" if m.group(2)[0].isdigit() else "List Bullet"
            p = doc.add_paragraph(style=style); add_inline(p, m.group(3)); i += 1; continue
        if not ln.strip():
            i += 1; continue
        # a paragraph: join following plain lines
        buf = [ln.strip()]; j = i + 1
        while j < len(lines) and lines[j].strip() and not re.match(r"^(#|```|\||\s*([-*]|\d+\.)\s)", lines[j]):
            buf.append(lines[j].strip()); j += 1
        p = doc.add_paragraph(); add_inline(p, " ".join(buf)); i = j
    doc.save(out)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    build(Path(sys.argv[1]).read_text(), sys.argv[2])
    print("wrote", sys.argv[2])
